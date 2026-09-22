import re
from pathlib import Path

from openpyxl import load_workbook

from app.core.phone_utils import normalize_korean_phone, format_korean_international


def normalize_phone(value):
    return normalize_korean_phone(value)


class DBImportService:
    def __init__(self, db, logs):
        self.db = db
        self.logs = logs

    def _create_import_run(self, source_file):
        return self.db.execute(
            "INSERT INTO import_runs(source_file) VALUES(?)",
            (source_file,),
        )

    def _import_values(self, values, source_file):
        import_id = self._create_import_run(source_file)

        stats = {
            "import_id": import_id,
            "total": 0,
            "valid": 0,
            "duplicate": 0,
            "invalid": 0,
            "duplicate_numbers": [],
        }
        seen_in_file = set()

        for raw in values:
            if raw is None:
                continue

            raw_text = str(raw).strip()
            if not raw_text:
                continue

            stats["total"] += 1
            normalized = normalize_phone(raw_text)

            if not normalized:
                stats["invalid"] += 1
                self.logs.write(
                    "WARNING",
                    "DB",
                    f"번호 형식 오류: {raw_text}",
                )
                continue

            reason = None
            if normalized in seen_in_file:
                reason = "업로드 파일 내부 중복"
            else:
                exists = self.db.fetchone(
                    "SELECT id FROM recipients WHERE normalized_phone=? LIMIT 1",
                    (normalized,),
                )
                if exists:
                    reason = "기존 데이터베이스 중복"

            if reason:
                stats["duplicate"] += 1
                stats["duplicate_numbers"].append(normalized)
                self.db.execute(
                    "INSERT INTO import_duplicates(import_id,raw_phone,normalized_phone,reason) "
                    "VALUES(?,?,?,?)",
                    (import_id, raw_text, normalized, reason),
                )
                self.logs.write(
                    "WARNING",
                    "DB",
                    f"중복번호 검수: {format_korean_international(normalized)} / {reason}",
                )
                continue

            seen_in_file.add(normalized)
            self.db.execute(
                "INSERT INTO recipients(source_file,import_id,phone,normalized_phone,status,contact_status) "
                "VALUES(?,?,?,?, 'PENDING','NOT_ADDED')",
                (source_file, import_id, raw_text, normalized),
            )
            stats["valid"] += 1

        self.db.execute(
            "UPDATE import_runs SET total_count=?,added_count=?,duplicate_count=?,invalid_count=? "
            "WHERE id=?",
            (
                stats["total"],
                stats["valid"],
                stats["duplicate"],
                stats["invalid"],
                import_id,
            ),
        )

        self.logs.write(
            "INFO",
            "DB",
            f"DB 업로드 완료 / 전체 {stats['total']} / 즉시사용 {stats['valid']} / "
            f"중복검수 {stats['duplicate']} / 오류 {stats['invalid']}",
        )
        return stats

    def import_xlsx(self, path):
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        source_file = Path(path).name

        values = []
        for row in ws.iter_rows(values_only=True):
            raw = row[0] if row else None
            if raw is not None:
                values.append(raw)

        try:
            wb.close()
        except Exception:
            pass

        return self._import_values(values, source_file)

    def import_txt(self, path):
        source_file = Path(path).name
        raw_bytes = Path(path).read_bytes()

        text = None
        for encoding in ("utf-8-sig", "utf-8", "cp949", "euc-kr"):
            try:
                text = raw_bytes.decode(encoding)
                break
            except UnicodeDecodeError:
                continue

        if text is None:
            raise ValueError("TXT 파일 인코딩을 읽을 수 없습니다. UTF-8 또는 CP949로 저장해주세요.")

        values = []
        for line in text.splitlines():
            line = line.strip()
            if not line:
                continue

            # 한 줄에 여러 번호가 있을 경우 쉼표/탭/세미콜론 기준 우선 분리
            parts = [p.strip() for p in re.split(r"[,;\t]+", line) if p.strip()]

            if len(parts) == 1:
                # 번호 하나가 공백을 포함한 국제형식(+82 10 1234 5678)일 수 있으므로
                # 먼저 전체 줄을 번호로 판단한다.
                if normalize_phone(line):
                    values.append(line)
                    continue

                # 전체 줄이 번호가 아니면 공백으로 분리된 토큰 중 전화번호 형태만 수집
                tokens = [p.strip() for p in re.split(r"\s+", line) if p.strip()]
                matched = False
                for token in tokens:
                    if normalize_phone(token):
                        values.append(token)
                        matched = True
                if not matched:
                    values.append(line)
            else:
                for part in parts:
                    if normalize_phone(part):
                        values.append(part)
                    else:
                        # "이름 010-1234-5678" 같은 형태도 번호 부분 추출
                        candidates = re.findall(
                            r"(?:\+?82[\s-]?)?0?1[016789](?:[\s-]?\d){7,8}",
                            part,
                        )
                        if candidates:
                            values.extend(candidates)
                        else:
                            values.append(part)

        return self._import_values(values, source_file)

    def import_file(self, path):
        suffix = Path(path).suffix.lower()
        if suffix == ".xlsx":
            return self.import_xlsx(path)
        if suffix == ".txt":
            return self.import_txt(path)
        raise ValueError("지원하지 않는 DB 파일입니다. XLSX 또는 TXT 파일을 선택해주세요.")

    def approve_duplicates(self, import_id):
        rows = self.db.fetchall(
            "SELECT * FROM import_duplicates WHERE import_id=? AND processed=0 ORDER BY id",
            (import_id,),
        )
        added = 0
        with self.db.connection() as conn:
            source = conn.execute(
                "SELECT source_file FROM import_runs WHERE id=?",
                (import_id,),
            ).fetchone()
            source_file = source["source_file"] if source else "중복검수"

            for row in rows:
                conn.execute(
                    "INSERT INTO recipients(source_file,import_id,phone,normalized_phone,status,contact_status) "
                    "VALUES(?,?,?,?, 'PENDING','NOT_ADDED')",
                    (
                        source_file,
                        import_id,
                        row["raw_phone"],
                        row["normalized_phone"],
                    ),
                )
                conn.execute(
                    "UPDATE import_duplicates SET approved=1,processed=1 WHERE id=?",
                    (row["id"],),
                )
                added += 1

            conn.execute(
                "UPDATE import_runs SET duplicate_decision='APPROVED',added_count=added_count+? "
                "WHERE id=?",
                (added, import_id),
            )

        self.logs.write("INFO", "DB", f"중복번호 {added}개를 사용 DB에 추가했습니다.")
        return added

    def reject_duplicates(self, import_id):
        count = self.db.execute_rowcount(
            "UPDATE import_duplicates SET approved=0,processed=1 "
            "WHERE import_id=? AND processed=0",
            (import_id,),
        )
        self.db.execute(
            "UPDATE import_runs SET duplicate_decision='REJECTED' WHERE id=?",
            (import_id,),
        )
        self.logs.write("INFO", "DB", f"중복번호 {count}개를 제외했습니다.")
        return count

    def get_duplicates(self, import_id):
        return self.db.fetchall(
            "SELECT normalized_phone,reason FROM import_duplicates "
            "WHERE import_id=? ORDER BY id",
            (import_id,),
        )


    def get_import_rows(self, import_id):
        rows = []

        normal = self.db.fetchall(
            "SELECT id,phone,normalized_phone,status FROM recipients "
            "WHERE import_id=? ORDER BY id",
            (import_id,),
        )
        for row in normal:
            rows.append({
                "phone": row["phone"],
                "normalized_phone": row["normalized_phone"],
                "type": "NORMAL",
                "reason": "",
            })

        duplicates = self.db.fetchall(
            "SELECT id,raw_phone,normalized_phone,reason,processed,approved "
            "FROM import_duplicates WHERE import_id=? ORDER BY id",
            (import_id,),
        )
        for row in duplicates:
            rows.append({
                "phone": row["raw_phone"],
                "normalized_phone": row["normalized_phone"],
                "type": "DUPLICATE",
                "reason": row["reason"],
            })

        return rows
