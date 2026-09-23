import csv
from pathlib import Path

from openpyxl import load_workbook, Workbook

from app.core.phone_utils import normalize_korean_phone
from app.core.paths import EXPORTS_DIR


MIN_COUNT = 100
MAX_COUNT = 1_000_000
UNIT_PRICE_TENTHS_KRW = 6


class TelegramCheckError(Exception):
    def __init__(self, code, message, data=None):
        super().__init__(message)
        self.code = code
        self.data = data or {}


def format_domestic_phone(value):
    digits = normalize_korean_phone(value)
    if not digits:
        return ""
    if digits.startswith("010") and len(digits) == 11:
        return f"{digits[:3]}-{digits[3:7]}-{digits[7:]}"
    if digits.startswith("02") and len(digits) in (9, 10):
        mid = 5 if len(digits) == 9 else 6
        return f"{digits[:2]}-{digits[2:mid]}-{digits[mid:]}"
    if len(digits) == 10:
        return f"{digits[:3]}-{digits[3:6]}-{digits[6:]}"
    if len(digits) == 11:
        return f"{digits[:3]}-{digits[3:7]}-{digits[7:]}"
    return digits

def to_api_phone(value):
    digits = normalize_korean_phone(value)
    if not digits:
        return ""
    return "82" + digits[1:]


class TelegramCheckService:
    def __init__(self, db, logs=None):
        self.db = db
        self.logs = logs

    def _iter_values(self, path):
        ext = Path(path).suffix.lower()

        if ext in {".txt", ".csv"}:
            opened = None
            for encoding in ("utf-8-sig", "utf-8", "cp949", "euc-kr"):
                try:
                    opened = open(path, "r", encoding=encoding, newline="")
                    opened.read(2048)
                    opened.seek(0)
                    break
                except UnicodeDecodeError:
                    if opened:
                        opened.close()
                    opened = None

            if opened is None:
                raise TelegramCheckError("FILE_ENCODING", "파일 인코딩을 읽을 수 없습니다.")

            with opened:
                if ext == ".txt":
                    for line in opened:
                        value = line.strip()
                        if value:
                            yield value
                else:
                    for row in csv.reader(opened):
                        if row and str(row[0]).strip():
                            yield str(row[0]).strip()
            return

        if ext == ".xlsx":
            wb = load_workbook(path, read_only=True, data_only=True)
            try:
                ws = wb.active
                for row in ws.iter_rows(values_only=True):
                    if row and row[0] is not None and str(row[0]).strip():
                        yield str(row[0]).strip()
            finally:
                wb.close()
            return

        raise TelegramCheckError(
            "FILE_TYPE",
            "지원하지 않는 파일입니다. TXT / CSV / XLSX만 업로드할 수 있습니다.",
        )

    def prepare_file(self, path, filter_type="ALL"):
        filter_type = str(filter_type or "ALL").upper()
        if filter_type not in {"1D", "3D", "7D", "ALL"}:
            raise TelegramCheckError("INVALID_FILTER", "검수 조건이 올바르지 않습니다.")

        task_id = self.db.execute(
            "INSERT INTO telegram_check_tasks(source_file,filter_type,status) "
            "VALUES(?,?,'DRAFT')",
            (Path(path).name, filter_type),
        )

        total = valid = duplicate = invalid = 0
        seen = set()
        batch = []

        with self.db.connection() as conn:
            for raw in self._iter_values(path):
                total += 1
                if total > MAX_COUNT:
                    conn.execute(
                        "UPDATE telegram_check_tasks SET status='FAILED',"
                        "error_code='MAX_COUNT',error_message=?,updated_at=CURRENT_TIMESTAMP "
                        "WHERE id=?",
                        ("1회 최대 검수 수량은 1,000,000건입니다.", task_id),
                    )
                    raise TelegramCheckError(
                        "MAX_COUNT",
                        "1회 최대 검수 수량은 1,000,000건입니다.",
                    )

                normalized = normalize_korean_phone(raw)
                if not normalized:
                    invalid += 1
                    batch.append((task_id, raw, None, None, "INVALID"))
                elif normalized in seen:
                    duplicate += 1
                    batch.append(
                        (task_id, raw, normalized, format_domestic_phone(normalized), "DUPLICATE")
                    )
                else:
                    seen.add(normalized)
                    valid += 1
                    batch.append(
                        (task_id, raw, normalized, format_domestic_phone(normalized), "VALID")
                    )

                if len(batch) >= 2000:
                    conn.executemany(
                        "INSERT INTO telegram_check_inputs("
                        "task_id,raw_phone,normalized_phone,display_phone,input_status"
                        ") VALUES(?,?,?,?,?)",
                        batch,
                    )
                    batch.clear()

            if batch:
                conn.executemany(
                    "INSERT INTO telegram_check_inputs("
                    "task_id,raw_phone,normalized_phone,display_phone,input_status"
                    ") VALUES(?,?,?,?,?)",
                    batch,
                )

            amount_tenths = valid * UNIT_PRICE_TENTHS_KRW
            conn.execute(
                "UPDATE telegram_check_tasks SET "
                "total_input_count=?,valid_count=?,duplicate_count=?,invalid_count=?,"
                "charged_count=?,amount_tenths_krw=?,updated_at=CURRENT_TIMESTAMP "
                "WHERE id=?",
                (total, valid, duplicate, invalid, valid, amount_tenths, task_id),
            )

        if self.logs:
            self.logs.write(
                "INFO",
                "가입자검수",
                f"파일 검증 완료 / 작업 #{task_id} / 입력 {total:,} / "
                f"정상 {valid:,} / 중복 {duplicate:,} / 오류 {invalid:,}",
            )

        return self.task_summary(task_id)

    def task_summary(self, task_id):
        row = self.db.fetchone("SELECT * FROM telegram_check_tasks WHERE id=?", (task_id,))
        if not row:
            raise TelegramCheckError("TASK_NOT_FOUND", "검수 작업을 찾을 수 없습니다.")
        data = dict(row)
        data["amount_krw"] = data["amount_tenths_krw"] / 10.0
        data["refund_krw"] = data["refund_tenths_krw"] / 10.0
        return data

    def list_tasks(self, limit=100):
        return self.db.fetchall(
            "SELECT * FROM telegram_check_tasks ORDER BY id DESC LIMIT ?",
            (int(limit),),
        )

    def can_start(self, task_id):
        row = self.task_summary(task_id)
        count = int(row["charged_count"] or 0)
        if count < MIN_COUNT:
            return False, "최소 검수 수량은 100건입니다."
        if count > MAX_COUNT:
            return False, "1회 최대 검수 수량은 1,000,000건입니다."
        return True, ""


    def export_api_input_txt(self, task_id, path):
        row = self.task_summary(task_id)
        ok, reason = self.can_start(task_id)
        if not ok:
            raise TelegramCheckError("COUNT_RANGE", reason)

        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("w", encoding="utf-8", newline="\n") as out:
            offset = 0
            chunk = 5000
            while True:
                rows = self.db.fetchall(
                    "SELECT normalized_phone FROM telegram_check_inputs "
                    "WHERE task_id=? AND input_status='VALID' "
                    "ORDER BY id LIMIT ? OFFSET ?",
                    (task_id, chunk, offset),
                )
                if not rows:
                    break
                for item in rows:
                    value = to_api_phone(item["normalized_phone"])
                    if value:
                        out.write(value + "\n")
                offset += len(rows)

        if self.logs:
            self.logs.write(
                "INFO", "가입자검수",
                f"검수 대상 TXT 저장 / 작업 #{task_id} / {int(row['charged_count']):,}건"
            )
        return str(path)

    def _iter_result_rows(self, path):
        opened = None
        for encoding in ("utf-8-sig", "utf-8", "gb18030", "cp949", "euc-kr"):
            try:
                opened = open(path, "r", encoding=encoding, newline="")
                opened.read(2048)
                opened.seek(0)
                break
            except UnicodeDecodeError:
                if opened:
                    opened.close()
                opened = None
        if opened is None:
            raise TelegramCheckError("RESULT_ENCODING", "결과 CSV 인코딩을 읽을 수 없습니다.")

        with opened:
            reader = csv.reader(opened)
            first = True
            for row in reader:
                if not row:
                    continue
                if first:
                    first = False
                    first_cell = str(row[0] or "").strip().lower()
                    if any(k in first_cell for k in ("phone", "전화", "mobile")):
                        continue
                yield row

    def import_result_csv(self, task_id, path):
        summary = self.task_summary(task_id)
        requested = int(summary["charged_count"] or 0)

        seen = set()
        rows_to_insert = []
        for row in self._iter_result_rows(path):
            raw_phone = row[0] if len(row) > 0 else ""
            normalized = normalize_korean_phone(raw_phone)
            if not normalized or normalized in seen:
                continue
            seen.add(normalized)

            telegram_id = str(row[1] or "").strip() if len(row) > 1 else ""
            telegram_username = str(row[2] or "").strip() if len(row) > 2 else ""
            telegram_active = str(row[-1] or "").strip() if len(row) > 3 else ""

            if telegram_id == "0":
                telegram_id = ""
            if telegram_username == "0":
                telegram_username = ""

            rows_to_insert.append((
                task_id,
                format_domestic_phone(normalized),
                "가입",
                telegram_id or None,
                telegram_username or None,
                telegram_active or None,
                "API_RESULT",
            ))

        with self.db.connection() as conn:
            conn.execute("DELETE FROM telegram_check_results WHERE task_id=?", (task_id,))
            if rows_to_insert:
                for start in range(0, len(rows_to_insert), 2000):
                    conn.executemany(
                        "INSERT INTO telegram_check_results("
                        "task_id,phone_number,telegram_check_status,telegram_id,"
                        "telegram_username,telegram_active,raw_status"
                        ") VALUES(?,?,?,?,?,?,?)",
                        rows_to_insert[start:start + 2000],
                    )

            success = len(rows_to_insert)
            missing = max(0, requested - success)
            status = "COMPLETED" if missing == 0 else "PARTIAL"
            conn.execute(
                "UPDATE telegram_check_tasks SET success_count=?,missing_count=?,"
                "status=?,completed_at=CURRENT_TIMESTAMP,updated_at=CURRENT_TIMESTAMP "
                "WHERE id=?",
                (success, missing, status, task_id),
            )

        if self.logs:
            self.logs.write(
                "SUCCESS", "가입자검수",
                f"결과 가져오기 완료 / 작업 #{task_id} / 결과 {len(rows_to_insert):,} / "
                f"누락 {max(0, requested-len(rows_to_insert)):,}"
            )
        return self.task_summary(task_id)

    def _active_limit(self, filter_type):
        return {"1D": 1, "3D": 3, "7D": 7}.get(str(filter_type or "").upper())

    def result_rows(self, task_id):
        summary = self.task_summary(task_id)
        limit = self._active_limit(summary["filter_type"])
        rows = self.db.fetchall(
            "SELECT phone_number,telegram_check_status,telegram_id,"
            "telegram_username,telegram_active "
            "FROM telegram_check_results WHERE task_id=? ORDER BY id",
            (task_id,),
        )
        result = []
        for row in rows:
            active_raw = row["telegram_active"]
            try:
                active_num = int(float(str(active_raw).strip()))
            except Exception:
                active_num = None
            if limit is not None and (active_num is None or active_num > limit):
                continue
            result.append({
                "전화번호": row["phone_number"] or "",
                "가입여부": row["telegram_check_status"] or "확인불가",
                "텔레그램 UID": row["telegram_id"] or "",
                "텔레그램 ID": row["telegram_username"] or "",
                "텔레그램 접속일자": "" if active_raw is None else str(active_raw),
            })
        result.sort(
            key=lambda x: (
                10**9 if not str(x["텔레그램 접속일자"]).strip().replace(".", "", 1).isdigit()
                else float(x["텔레그램 접속일자"])
            )
        )
        return result

    def export_result_csv(self, task_id, path):
        rows = self.result_rows(task_id)
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        columns = ["전화번호", "가입여부", "텔레그램 UID", "텔레그램 ID", "텔레그램 접속일자"]
        with path.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=columns)
            writer.writeheader()
            writer.writerows(rows)
        self.db.execute(
            "UPDATE telegram_check_tasks SET result_file_path=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",
            (str(path), task_id),
        )
        return str(path), len(rows)

    def export_result_xlsx(self, task_id, path):
        rows = self.result_rows(task_id)
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("검수결과")
        columns = ["전화번호", "가입여부", "텔레그램 UID", "텔레그램 ID", "텔레그램 접속일자"]
        ws.append(columns)
        for row in rows:
            ws.append([row[col] for col in columns])
        wb.save(path)
        self.db.execute(
            "UPDATE telegram_check_tasks SET result_file_path=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",
            (str(path), task_id),
        )
        return str(path), len(rows)

    def cancel_task(self, task_id):
        row = self.task_summary(task_id)
        if row["status"] in ("COMPLETED", "PARTIAL"):
            raise TelegramCheckError("TASK_FINISHED", "이미 완료된 작업은 취소할 수 없습니다.")
        self.db.execute(
            "UPDATE telegram_check_tasks SET status='CANCELLED',updated_at=CURRENT_TIMESTAMP WHERE id=?",
            (task_id,),
        )
        return self.task_summary(task_id)
