from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from app.core.paths import EXPORTS_DIR


class CompletionExportService:
    def __init__(self, db, logs):
        self.db = db
        self.logs = logs

    def completion_rows(self, campaign_id):
        rows = self.db.fetchall(
            "SELECT "
            "cr.recipient_id, r.phone, r.normalized_phone, r.contact_name, "
            "cr.contact_status, cr.contact_added_at, cr.status AS send_status, "
            "cr.sent_at, cr.postbot_code, cr.telegram_message_id, "
            "cr.error_code, cr.error_message, "
            "a.name AS account_name, a.phone AS account_phone "
            "FROM campaign_recipients cr "
            "JOIN recipients r ON r.id=cr.recipient_id "
            "LEFT JOIN telegram_accounts a ON a.id=cr.assigned_account_id "
            "WHERE cr.campaign_id=? "
            "ORDER BY cr.recipient_id",
            (campaign_id,),
        )

        result = []
        for index, row in enumerate(rows, start=1):
            contact_ok = row["contact_status"] == "ADDED"
            send_ok = row["send_status"] == "MESSAGE_SENT"

            result.append({
                "db_no": f"DB{index}",
                "recipient_id": row["recipient_id"],
                "phone": row["phone"] or "",
                "normalized_phone": row["normalized_phone"] or "",
                "contact_name": row["contact_name"] or "",
                "contact_status": "연락처 추가 성공" if contact_ok else f"연락처 {row['contact_status']}",
                "contact_time": row["contact_added_at"] or "",
                "send_status": "메시지 전송 성공" if send_ok else f"메시지 {row['send_status']}",
                "send_time": row["sent_at"] or "",
                "postbot_code": row["postbot_code"] or "",
                "account": row["account_phone"] or row["account_name"] or "",
                "message_id": row["telegram_message_id"] or "",
                "error": " / ".join(
                    x for x in [row["error_code"] or "", row["error_message"] or ""] if x
                ),
            })

        return result

    def export_xlsx(self, campaign_id, target_path=None):
        campaign = self.db.fetchone(
            "SELECT * FROM campaigns WHERE id=?",
            (campaign_id,),
        )
        if not campaign:
            raise RuntimeError("작업완료 로그를 찾을 수 없습니다.")

        rows = self.completion_rows(campaign_id)

        if target_path:
            output = Path(target_path)
        else:
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output = EXPORTS_DIR / f"AngelToggle_작업완료_{campaign_id}_{stamp}.xlsx"

        output.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "작업완료"

        headers = [
            "DB번호",
            "원본 전화번호",
            "변환 전화번호",
            "연락처 저장명",
            "연락처 추가 상태",
            "연락처 추가 성공시간",
            "메시지 전송 상태",
            "메시지 성공시간",
            "PostBot 고유번호",
            "사용 계정",
            "Telegram Message ID",
            "오류",
        ]
        ws.append(headers)

        for row in rows:
            ws.append([
                row["db_no"],
                row["phone"],
                row["normalized_phone"],
                row["contact_name"],
                row["contact_status"],
                row["contact_time"],
                row["send_status"],
                row["send_time"],
                row["postbot_code"],
                row["account"],
                row["message_id"],
                row["error"],
            ])

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        widths = [10, 18, 18, 24, 20, 22, 20, 22, 24, 18, 22, 36]
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        wb.save(output)

        self.logs.write(
            "INFO",
            "작업완료",
            f"작업 #{campaign_id} 완료 DB Excel 저장 / {output}",
            campaign_id=campaign_id,
        )
        return str(output)
